"""tools.py — utility: genera INDICE.md ed esegue un controllo di overflow.

Uso:  python3 docs/presentazione/tools.py
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "Delta5_iCE2026_6_CELL_Training.pptx")
OUT = os.path.join(HERE, "INDICE.md")
EMU_IN = 914400
IN_PT = 72


def build_index(prs):
    def title_of(slide):
        cands = []
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.bold and r.font.size.pt in (24, 28, 30, 36):
                        cands.append((r.font.size.pt, p.text))
        cands.sort(key=lambda c: -c[0])
        return cands[0][1] if cands else ""

    lines = ["# Indice della presentazione", "",
             "`Delta5_iCE2026_6_CELL_Training.pptx` — %d slide 16:9, note del relatore "
             "su ogni slide." % len(prs.slides), ""]
    for i, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        big = any(r.font.size and r.font.size.pt >= 70
                  for sh in slide.shapes if sh.has_text_frame
                  for p in sh.text_frame.paragraphs for r in p.runs)
        if big:
            lines.append("\n## %s · %s   (da slide %d)"
                         % (texts[0], texts[1] if len(texts) > 1 else "", i))
        else:
            lines.append("- %d. %s" % (i, "Copertina" if i == 1 else title_of(slide)))
    return "\n".join(lines) + "\n"


def est_lines(text, width_pt, size_pt, mono=False):
    if not text:
        return 0
    cpl = max(8, int(width_pt / (size_pt * (0.60 if mono else 0.50))))
    return sum(max(1, -(-len(line) // cpl)) for line in text.split("\n"))


def check(prs):
    problems = []
    for idx, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                tf = sh.text_frame
                if not tf.text.strip():
                    continue
                w = sh.width / EMU_IN * IN_PT
                h = sh.height / EMU_IN * IN_PT
                if h < 110:      # titoli, kicker, footer
                    continue
                tot = 0
                for p in tf.paragraphs:
                    size, mono = 16, False
                    for r in p.runs:
                        if r.font.size:
                            size = r.font.size.pt
                        if r.font.name and "Consolas" in r.font.name:
                            mono = True
                    lvl = p.level or 0
                    tot += est_lines(p.text, w - lvl * 18, size, mono) * size * 1.22
                    tot += (p.space_after.pt if p.space_after else 0)
                if tot > h * 1.06 + 4:
                    problems.append((idx, int(tot), int(h), tf.text[:50]))
            if sh.has_table:
                need = 0
                for row in sh.table.rows:
                    rowmax = 0
                    for i, cell in enumerate(row.cells):
                        size = 12
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size:
                                    size = r.font.size.pt
                        cw = sh.table.columns[i].width / EMU_IN * IN_PT - 10
                        rowmax = max(rowmax, est_lines(cell.text, cw, size) * size * 1.25 + 4)
                    need += rowmax
                avail = sh.height / EMU_IN * IN_PT
                if need > avail * 1.10 + 8:
                    problems.append((idx, int(need), int(avail), "TABLE"))
    return problems


def main():
    from pptx import Presentation
    prs = Presentation(PPTX)
    with open(OUT, "w") as fh:
        fh.write(build_index(prs))
    print("indice scritto:", OUT)
    probs = check(prs)
    print("slide:", len(prs.slides), "· possibili overflow:", len(probs))
    for p in probs:
        print("   s%3d stimato=%5dpt disponibile=%5dpt  %s" % p)


if __name__ == "__main__":
    main()
