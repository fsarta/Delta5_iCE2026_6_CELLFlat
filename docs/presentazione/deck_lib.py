"""deck_lib.py — motore di impaginazione per la presentazione di training
Delta5_iCE2026_6_CELLFlat (16:9, note del relatore, header/footer)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x23, 0x39, 0x4A)
STEEL = RGBColor(0x3D, 0x61, 0x7E)
ACCENT = RGBColor(0x2E, 0x7D, 0xA0)
LIGHT = RGBColor(0xE8, 0xED, 0xF1)
PAPER = RGBColor(0xF7, 0xF9, 0xFA)
GREY = RGBColor(0x5D, 0x6E, 0x7C)
DARKGREY = RGBColor(0x1F, 0x2B, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODEBG = RGBColor(0xF2, 0xF5, 0xF7)
MUTED = RGBColor(0x9F, 0xC4, 0xD6)
SOFT = RGBColor(0xC9, 0xDA, 0xE4)

FONT = "Calibri"
MONO = "Consolas"
W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self, project):
        self.prs = Presentation()
        self.prs.slide_width = W
        self.prs.slide_height = H
        self.project = project
        self.section_name = "Introduzione"
        self.n = 0
        self.blank = self.prs.slide_layouts[6]

    # ---------------------------------------------------------- primitive
    def _slide(self):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        return s

    def _rect(self, s, x, y, w, h, fill, line=None):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(0.75)
        sh.shadow.inherit = False
        return sh

    def _text(self, s, x, y, w, h, text, size=16, bold=False, color=DARKGREY,
              align=PP_ALIGN.LEFT, font=FONT, spacing=1.0, space_after=4):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            p.space_after = Pt(space_after)
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = color
        return tb

    def _bullets(self, s, x, y, w, h, items, size=16):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        first = True
        for it in items:
            lvl, txt = it if isinstance(it, (tuple, list)) else (0, it)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = min(lvl, 4)
            p.space_after = Pt(5 if lvl == 0 else 2)
            p.line_spacing = 1.0
            r = p.add_run()
            r.text = ({0: "", 1: "–  ", 2: "·  "}.get(lvl, "·  ")) + txt
            r.font.size = Pt(size - 1.5 * lvl)
            r.font.name = FONT
            r.font.color.rgb = DARKGREY if lvl == 0 else GREY
        return tb

    def _chrome(self, s, title, kicker=None):
        self._rect(s, 0, 0, W, Inches(0.92), NAVY)
        self._rect(s, 0, Inches(0.92), W, Inches(0.055), ACCENT)
        self._text(s, Inches(0.55), Inches(0.14), Inches(11.2), Inches(0.34),
                   (kicker or self.section_name).upper(), size=10.5, bold=True,
                   color=MUTED)
        self._text(s, Inches(0.55), Inches(0.42), Inches(11.8), Inches(0.45),
                   title, size=24, bold=True, color=WHITE)
        self._text(s, Inches(0.55), Inches(7.08), Inches(8.0), Inches(0.3),
                   self.project + "  ·  training tecnico", size=8.5, color=GREY)
        self._text(s, Inches(11.6), Inches(7.08), Inches(1.2), Inches(0.3),
                   str(self.n), size=8.5, color=GREY, align=PP_ALIGN.RIGHT)

    def _notes(self, s, notes):
        if not notes:
            return
        tf = s.notes_slide.notes_text_frame
        tf.text = notes
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(12)

    # ---------------------------------------------------------- tipi di slide
    def cover(self, title, subtitle, meta, notes=None):
        s = self._slide()
        self._rect(s, 0, 0, W, H, NAVY)
        self._rect(s, 0, Inches(4.62), W, Inches(0.06), ACCENT)
        self._text(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.4),
                   "FORMAZIONE TECNICA INTERNA", size=13, bold=True, color=MUTED)
        self._text(s, Inches(0.9), Inches(2.0), Inches(11.6), Inches(1.7),
                   title, size=44, bold=True, color=WHITE, spacing=0.95)
        self._text(s, Inches(0.9), Inches(3.72), Inches(11.0), Inches(0.8),
                   subtitle, size=18, color=SOFT)
        y = Inches(4.95)
        for line in meta:
            self._text(s, Inches(0.9), y, Inches(11.4), Inches(0.32),
                       line, size=12.5, color=MUTED)
            y += Inches(0.36)
        self._notes(s, notes)

    def section(self, num, title, subtitle="", notes="", topics=None):
        s = self._slide()
        self.section_name = title
        self._rect(s, 0, 0, W, H, STEEL)
        self._rect(s, 0, 0, Inches(0.34), H, ACCENT)
        self._text(s, Inches(1.0), Inches(1.5), Inches(3.0), Inches(1.2),
                   num, size=72, bold=True, color=RGBColor(0x7F, 0xA5, 0xBE))
        self._text(s, Inches(1.0), Inches(2.8), Inches(11.0), Inches(1.1),
                   title, size=36, bold=True, color=WHITE, spacing=0.95)
        if subtitle:
            self._text(s, Inches(1.0), Inches(4.0), Inches(10.6), Inches(0.9),
                       subtitle, size=16, color=SOFT)
        if topics:
            y = Inches(5.0)
            for t in topics:
                self._text(s, Inches(1.0), y, Inches(11.0), Inches(0.3),
                           "—  " + t, size=12.5, color=RGBColor(0xB8, 0xCD, 0xDB))
                y += Inches(0.32)
        self._notes(s, notes)

    def bullets(self, title, items, notes=None, kicker=None, size=16, footer=None):
        s = self._slide()
        self._chrome(s, title, kicker)
        self._bullets(s, Inches(0.55), Inches(1.28), Inches(12.2), Inches(5.4),
                      items, size=size)
        if footer:
            self._rect(s, Inches(0.55), Inches(6.52), Inches(12.2), Inches(0.02), LIGHT)
            self._text(s, Inches(0.55), Inches(6.62), Inches(12.2), Inches(0.35),
                       footer, size=10.5, color=ACCENT, bold=True)
        self._notes(s, notes)

    def two(self, title, lt, li, rt, ri, notes=None, kicker=None, size=12):
        s = self._slide()
        self._chrome(s, title, kicker)
        self._rect(s, Inches(0.55), Inches(1.3), Inches(6.0), Inches(0.42), LIGHT)
        self._text(s, Inches(0.68), Inches(1.36), Inches(5.8), Inches(0.35),
                   lt, size=13.5, bold=True, color=NAVY)
        self._bullets(s, Inches(0.62), Inches(1.88), Inches(5.9), Inches(4.9), li, size=size)
        self._rect(s, Inches(6.78), Inches(1.3), Inches(6.0), Inches(0.42), LIGHT)
        self._text(s, Inches(6.91), Inches(1.36), Inches(5.8), Inches(0.35),
                   rt, size=13.5, bold=True, color=NAVY)
        self._bullets(s, Inches(6.85), Inches(1.88), Inches(5.9), Inches(4.9), ri, size=size)
        self._notes(s, notes)

    def table(self, title, headers, rows, notes=None, kicker=None,
              col_widths=None, size=11, footer=None):
        s = self._slide()
        self._chrome(s, title, kicker)
        shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.55),
                                   Inches(1.3), Inches(12.2),
                                   Inches(5.0 if footer else 5.6))
        tbl = shape.table
        if col_widths:
            tot = sum(col_widths)
            for i, cw in enumerate(col_widths):
                tbl.columns[i].width = Emu(int(Inches(12.2) * cw / tot))
        for j, htxt in enumerate(headers):
            c = tbl.cell(0, j)
            c.text = ""
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY
            r = c.text_frame.paragraphs[0].add_run()
            r.text = htxt
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.name = FONT
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                c = tbl.cell(i, j)
                c.text = ""
                c.fill.solid()
                c.fill.fore_color.rgb = WHITE if i % 2 else PAPER
                c.margin_left = Inches(0.06)
                c.margin_right = Inches(0.06)
                c.margin_top = Inches(0.02)
                c.margin_bottom = Inches(0.02)
                txt = str(val)
                bold = txt.startswith("**")
                if bold:
                    txt = txt.strip("*")
                r = c.text_frame.paragraphs[0].add_run()
                r.text = txt
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.name = FONT
                r.font.color.rgb = DARKGREY
        if footer:
            self._text(s, Inches(0.55), Inches(6.62), Inches(12.2), Inches(0.35),
                       footer, size=10.5, color=ACCENT, bold=True)
        self._notes(s, notes)

    def code(self, title, code_text, caption=None, notes=None, kicker=None, size=11):
        s = self._slide()
        self._chrome(s, title, kicker)
        h = Inches(5.3) if caption else Inches(5.65)
        self._rect(s, Inches(0.55), Inches(1.28), Inches(12.2), h, CODEBG, line=LIGHT)
        self._rect(s, Inches(0.55), Inches(1.28), Inches(0.05), h, ACCENT)
        self._text(s, Inches(0.78), Inches(1.40), Inches(11.8), h - Inches(0.2),
                   code_text, size=size, font=MONO,
                   color=RGBColor(0x1B, 0x2B, 0x38), spacing=1.05, space_after=0)
        if caption:
            self._text(s, Inches(0.55), Inches(6.62), Inches(12.2), Inches(0.4),
                       caption, size=10.5, color=ACCENT, bold=True)
        self._notes(s, notes)

    def statement(self, big, small="", notes=None):
        s = self._slide()
        self._rect(s, 0, 0, W, H, PAPER)
        self._rect(s, 0, Inches(3.2), Inches(0.16), Inches(1.4), ACCENT)
        self._text(s, Inches(1.0), Inches(3.05), Inches(11.0), Inches(1.4),
                   big, size=28, bold=True, color=NAVY, spacing=0.95)
        if small:
            self._text(s, Inches(1.0), Inches(4.6), Inches(10.6), Inches(1.2),
                       small, size=15, color=GREY)
        self._text(s, Inches(11.6), Inches(7.08), Inches(1.2), Inches(0.3),
                   str(self.n), size=8.5, color=GREY, align=PP_ALIGN.RIGHT)
        self._notes(s, notes)

    def save(self, path):
        self.prs.save(path)
        return self.n
